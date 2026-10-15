"""Bulk user import and its template — port of bulk_import_users_view / download_user_template_view."""
import io
import random
import string
from typing import Any

import pandas as pd
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.services.org_common import assign_default_role
from app.core import email as core_email
from app.models import User, UserType
from app.services.users_common import create_user, msg

REQUIRED_COLUMNS = ["username", "email", "first_name", "last_name", "user_type"]
ALLOWED_EXTENSIONS = ["csv", "xlsx", "xls"]


def validate_file_name(filename: str | None) -> str | None:
    """forms.FileField(required) + FileExtensionValidator(['csv','xlsx','xls']) error text."""
    if not filename:
        return "This field is required."
    ext = filename.rsplit(".", 1)[-1] if "." in filename else ""
    if ext.lower() not in ALLOWED_EXTENSIONS:
        return (f"File extension “{ext}” is not allowed. Allowed extensions are: "
                f"{', '.join(ALLOWED_EXTENSIONS)}.")
    return None


def _send_welcome(db: Session, user: User, password: str) -> None:
    """send_mail(..., fail_silently=True) wrapped in try/except: never raises."""
    try:
        core_email.deliver(
            db, "Welcome to TuraVector",
            f"Hello {user.first_name},\n\n"
            f"Your account has been created successfully.\n\n"
            f"Username: {user.username}\n"
            f"Password: {password}\n\n"
            f"Please login and change your password immediately.\n\n"
            f"Best regards,\nTuraVector Team",
            "noreply@turadome.com", user.email,
        )
    except Exception:  # noqa: BLE001
        pass


def _find_user_type(db: Session, value) -> UserType:
    rows = db.execute(select(UserType).where(
        or_(func.lower(UserType.name) == str(value).lower(), func.lower(UserType.code) == str(value).lower()),
        UserType.is_active == 1,
    )).scalars().all()
    if not rows:
        raise LookupError
    if len(rows) > 1:
        raise Exception(f"get() returned more than one UserType -- it returned {len(rows)}!")
    return rows[0]


class _RowError(Exception):
    pass


def import_users(db: Session, filename: str, content: bytes, send_welcome_email: bool, skip_errors: bool,
                 allow_managers: bool = True, default_roles: bool = False) -> dict[str, Any]:
    """Returns {success, messages, redirect}. `redirect` is 'manage_users' or 'bulk_import_users' like Django."""
    try:
        if filename.endswith(".csv"):
            df = pd.read_csv(io.BytesIO(content))
        else:
            df = pd.read_excel(io.BytesIO(content))

        missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
        if missing:
            return {"success": False, "redirect": "bulk_import_users",
                    "messages": [msg("error", f"Missing required columns: {', '.join(missing)}")]}

        success_count = 0
        error_count = 0
        errors: list[str] = []

        outer = db.begin_nested()  # transaction.atomic()
        try:
            for index, row in df.iterrows():
                row_sp = db.begin_nested()
                try:
                    if pd.isna(row["username"]) or pd.isna(row["email"]):
                        row_sp.commit()
                        continue

                    if db.execute(select(User.id).where(User.username == row["username"]).limit(1)).first():
                        error_msg = f'Row {index + 2}: Username "{row["username"]}" already exists'
                        errors.append(error_msg)
                        if not skip_errors:
                            raise _RowError(error_msg)
                        error_count += 1
                        row_sp.commit()
                        continue

                    if db.execute(select(User.id).where(User.email == row["email"]).limit(1)).first():
                        error_msg = f'Row {index + 2}: Email "{row["email"]}" already exists'
                        errors.append(error_msg)
                        if not skip_errors:
                            raise _RowError(error_msg)
                        error_count += 1
                        row_sp.commit()
                        continue

                    try:
                        user_type = _find_user_type(db, row["user_type"])
                    except LookupError:
                        error_msg = f'Row {index + 2}: Invalid user type "{row["user_type"]}"'
                        errors.append(error_msg)
                        if not skip_errors:
                            raise _RowError(error_msg)
                        error_count += 1
                        row_sp.commit()
                        continue

                    if user_type.code == "manager" and not allow_managers:
                        error_msg = f"Row {index + 2}: Only a superadmin can create managers"
                        errors.append(error_msg)
                        if not skip_errors:
                            raise _RowError(error_msg)
                        error_count += 1
                        row_sp.commit()
                        continue

                    password = row.get("password", "") if "password" in df.columns else ""
                    if pd.isna(password) or not password:
                        password = "".join(random.choices(string.ascii_letters + string.digits, k=10))
                    password = str(password)

                    user = create_user(
                        db, username=str(row["username"]), email=str(row["email"]), password=password,
                        first_name=str(row["first_name"]), last_name=str(row["last_name"]),
                    )
                    user.user_type_new_id = user_type.id
                    user.user_type = user_type.code
                    if "phone" in df.columns and not pd.isna(row.get("phone")):
                        user.phone = str(row["phone"])
                    if "address" in df.columns and not pd.isna(row.get("address")):
                        user.address = str(row["address"])
                    if "is_manager" in df.columns and allow_managers:
                        user.is_manager = int(str(row.get("is_manager", "")).lower() in ["true", "yes", "1", "on"])
                    if default_roles:
                        assign_default_role(db, user)
                    db.flush()
                    row_sp.commit()

                    if send_welcome_email:
                        _send_welcome(db, user, password)
                    success_count += 1
                except Exception as e:  # noqa: BLE001
                    if row_sp.is_active:
                        row_sp.rollback()
                    if not skip_errors:
                        raise
                    errors.append(f"Row {index + 2}: {str(e)}")
                    error_count += 1
            outer.commit()
        except Exception:
            if outer.is_active:
                outer.rollback()
            raise
        db.commit()

        messages = []
        if success_count > 0:
            messages.append(msg("success", f"Successfully imported {success_count} users."))
        if error_count > 0:
            messages.append(msg("warning", f"Failed to import {error_count} users."))
            for error in errors[:5]:
                messages.append(msg("error", error))
            if len(errors) > 5:
                messages.append(msg("error", f"... and {len(errors) - 5} more errors."))
        return {"success": True, "redirect": "manage_users", "messages": messages}
    except Exception as e:  # noqa: BLE001
        return {"success": False, "redirect": "bulk_import_users",
                "messages": [msg("error", f"Error processing file: {str(e)}")]}


def build_template(db: Session, file_format: str) -> tuple[bytes, str, str]:
    """Returns (content, media_type, filename)."""
    user_types = list(db.execute(select(UserType.name).where(UserType.is_active == 1).order_by(UserType.name)).scalars())
    user_types_str = "/".join(user_types)
    first = user_types[0] if user_types else "Student"
    sample_data = {
        "username": ["john_doe", "jane_smith", "bob_jones"],
        "email": ["john@example.com", "jane@example.com", "bob@example.com"],
        "first_name": ["John", "Jane", "Bob"],
        "last_name": ["Doe", "Smith", "Jones"],
        "user_type": [first, first, first],
        "password": ["", "", ""],
        "phone": ["+1234567890", "+0987654321", ""],
        "address": ["123 Main St, City", "456 Oak Ave, Town", ""],
        "is_manager": ["false", "false", "true"],
    }
    df = pd.DataFrame(sample_data)
    instructions = {
        "username": "REQUIRED: Unique username",
        "email": "REQUIRED: Unique email",
        "first_name": "REQUIRED: First name",
        "last_name": "REQUIRED: Last name",
        "user_type": f"REQUIRED: One of {user_types_str}",
        "password": "OPTIONAL: Leave empty to auto-generate",
        "phone": "OPTIONAL: Phone number",
        "address": "OPTIONAL: Full address",
        "is_manager": "OPTIONAL: true/false",
    }
    df_final = pd.concat([pd.DataFrame([instructions]), df], ignore_index=True)

    if file_format == "xlsx":
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df_final.to_excel(writer, sheet_name="Users", index=False)
            worksheet = writer.sheets["Users"]
            for col in range(1, len(df_final.columns) + 1):
                cell = worksheet.cell(row=1, column=col)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                cell.font = Font(color="FFFFFF", bold=True)
            for col in range(1, len(df_final.columns) + 1):
                cell = worksheet.cell(row=2, column=col)
                cell.font = Font(italic=True, color="FF0000")
                cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "user_import_template.xlsx"

    output = io.StringIO()
    df_final.to_csv(output, index=False)
    return output.getvalue().encode(), "text/csv", "user_import_template.csv"
