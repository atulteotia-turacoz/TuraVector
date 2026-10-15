"""Report exports — line-for-line port of library.views.export_excel_report / export_pdf_report /
export_section_view helpers. Same sheet titles, headers, styles, column widths, PDF tables and file names."""
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.utils import user_type_display
from app.models import utcnow

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _styles():
    return (
        Font(name="Arial", size=16, bold=True),
        Font(name="Arial", size=12, bold=True, color="FFFFFF"),
        Font(name="Arial", size=10),
        PatternFill(start_color="366092", end_color="366092", fill_type="solid"),
        Alignment(horizontal="center", vertical="center"),
    )


def _autowidth(ws) -> None:
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:  # noqa: BLE001
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _slug(period_name: str) -> str:
    return period_name.lower().replace(" ", "_")


def _name(user) -> str:
    return f"{user.first_name} {user.last_name}"


def _header_table_style(header_size: int, body_size: int | None = None) -> TableStyle:
    cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.darkblue),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), header_size),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
        ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
    ]
    if body_size:
        cmds.append(("FONTSIZE", (0, 1), (-1, -1), body_size))
    return TableStyle(cmds)


# ------------------------------------------------------------------ full report
def excel_report(report: dict[str, Any]) -> tuple[bytes, str]:
    wb = Workbook()
    ws = wb.active
    ws.title = "TuraVector Report"
    title_font, header_font, regular_font, header_fill, center = _styles()

    ws.merge_cells("A1:F1")
    ws["A1"] = "TuraVector Analytics Report"
    ws["A1"].font = title_font
    ws["A1"].alignment = center

    ws.merge_cells("A2:F2")
    ws["A2"] = (f"Period: {report['period']} ({report['start_date'].strftime('%Y-%m-%d')} to "
                f"{report['end_date'].strftime('%Y-%m-%d')})")
    ws["A2"].font = regular_font
    ws["A2"].alignment = center

    current_row = 4
    ws.merge_cells(f"A{current_row}:F{current_row}")
    ws[f"A{current_row}"] = "Key Metrics"
    ws[f"A{current_row}"].font = header_font
    ws[f"A{current_row}"].fill = header_fill
    ws[f"A{current_row}"].alignment = center

    current_row += 1
    metrics_data = [
        ["Metric", "Total", "In Period"],
        ["Documents", report["total_documents"], report["documents_in_period"]],
        ["Versions", report["total_versions"], report["versions_in_period"]],
        ["References", report["total_references"], report["references_in_period"]],
        ["Active Users", "-", report["active_users"]],
    ]
    for row_data in metrics_data:
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col)
            cell.value = value
            if current_row == 5:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
            else:
                cell.font = regular_font
                cell.alignment = center
        current_row += 1

    current_row += 1
    ws.merge_cells(f"A{current_row}:F{current_row}")
    ws[f"A{current_row}"] = "Document Status Breakdown"
    ws[f"A{current_row}"].font = header_font
    ws[f"A{current_row}"].fill = header_fill
    ws[f"A{current_row}"].alignment = center

    current_row += 1
    status_data = [
        ["Status", "Count"],
        ["Draft", report["draft_docs"]],
        ["Under Review", report["review_docs"]],
        ["Approved", report["approved_docs"]],
        ["Revised", report["revised_docs"]],
    ]
    for row_data in status_data:
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col)
            cell.value = value
            # Django's header test here never matches, so every row gets the regular style.
            if current_row == (current_row - len(status_data) + 1):
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
            else:
                cell.font = regular_font
                cell.alignment = center
        current_row += 1

    current_row += 1
    ws.merge_cells(f"A{current_row}:F{current_row}")
    ws[f"A{current_row}"] = "Recent Documents"
    ws[f"A{current_row}"].font = header_font
    ws[f"A{current_row}"].fill = header_fill
    ws[f"A{current_row}"].alignment = center

    current_row += 1
    for col, header in enumerate(["Document Name", "Document Number", "Owner", "Status", "Created", "Type"], 1):
        cell = ws.cell(row=current_row, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    current_row += 1
    status_map = {0: "Draft", 1: "Under Review", 2: "Approved", 3: "Revised"}
    for doc in report["recent_documents"]:
        doc_data = [
            doc.document_name,
            doc.document_number,
            _name(doc.user),
            status_map.get(doc.stage_status, "Unknown"),
            doc.created_at.strftime("%Y-%m-%d"),
            doc.document_type.name if doc.document_type else "N/A",
        ]
        for col, value in enumerate(doc_data, 1):
            cell = ws.cell(row=current_row, column=col)
            cell.value = value
            cell.font = regular_font
        current_row += 1

    _autowidth(ws)
    return _save(wb), f"turavector_report_{_slug(report['period'])}.xlsx"


def pdf_report(report: dict[str, Any]) -> tuple[bytes, str]:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("CustomTitle", parent=styles["Heading1"], fontSize=18, spaceAfter=30, alignment=TA_CENTER)
    heading_style = ParagraphStyle("CustomHeading", parent=styles["Heading2"], fontSize=14, spaceAfter=12,
                                   textColor=colors.darkblue)
    story = [
        Paragraph("TuraVector Analytics Report", title_style),
        Paragraph(f"Period: {report['period']} ({report['start_date'].strftime('%Y-%m-%d')} to "
                  f"{report['end_date'].strftime('%Y-%m-%d')})", styles["Normal"]),
        Spacer(1, 20),
        Paragraph("Key Metrics", heading_style),
    ]
    metrics_data = [
        ["Metric", "Total", "In Period"],
        ["Documents", str(report["total_documents"]), str(report["documents_in_period"])],
        ["Versions", str(report["total_versions"]), str(report["versions_in_period"])],
        ["References", str(report["total_references"]), str(report["references_in_period"])],
        ["Active Users", "-", str(report["active_users"])],
    ]
    t = Table(metrics_data, colWidths=[2 * inch, 1 * inch, 1 * inch])
    t.setStyle(_header_table_style(12))
    story += [t, Spacer(1, 20), Paragraph("Document Status Breakdown", heading_style)]

    status_data = [
        ["Status", "Count"],
        ["Draft", str(report["draft_docs"])],
        ["Under Review", str(report["review_docs"])],
        ["Approved", str(report["approved_docs"])],
        ["Revised", str(report["revised_docs"])],
    ]
    t = Table(status_data, colWidths=[2 * inch, 1 * inch])
    t.setStyle(_header_table_style(12))
    story += [t, Spacer(1, 20), Paragraph("Recent Documents", heading_style)]

    doc_data = [["Document Name", "Owner", "Status", "Created"]]
    status_map = {0: "Draft", 1: "Review", 2: "Approved", 3: "Revised"}
    for d in report["recent_documents"][:5]:
        doc_data.append([
            d.document_name[:30] + "..." if len(d.document_name) > 30 else d.document_name,
            _name(d.user),
            status_map.get(d.stage_status, "Unknown"),
            d.created_at.strftime("%Y-%m-%d"),
        ])
    t = Table(doc_data, colWidths=[2.5 * inch, 1.5 * inch, 1 * inch, 1 * inch])
    t.setStyle(_header_table_style(10, 9))
    story.append(t)
    doc.build(story)
    return buffer.getvalue(), f"turavector_report_{_slug(report['period'])}.pdf"


# ------------------------------------------------------------------ sections
def _section_sheet(title: str, sheet: str, merge_to: str, period_name: str, headers: list[str]):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    title_font, header_font, regular_font, header_fill, center = _styles()
    ws.merge_cells(f"A1:{merge_to}1")
    ws["A1"] = title
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws.merge_cells(f"A2:{merge_to}2")
    ws["A2"] = f"Period: {period_name} (as of {utcnow().strftime('%Y-%m-%d %H:%M')})"
    ws["A2"].font = regular_font
    ws["A2"].alignment = center
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
    return wb, ws


def _section_pdf(title: str, period_name: str, data: list[list[str]], col_widths: list[float]) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("CustomTitle", parent=styles["Heading1"], fontSize=18, spaceAfter=30, alignment=TA_CENTER)
    table = Table(data, colWidths=[w * inch for w in col_widths])
    table.setStyle(_header_table_style(10, 8))
    doc.build([
        Paragraph(title, title_style),
        Paragraph(f"Period: {period_name} (as of {utcnow().strftime('%Y-%m-%d %H:%M')})", styles["Normal"]),
        Spacer(1, 20),
        table,
    ])
    return buffer.getvalue()


def recent_documents_excel(documents, period_name: str) -> tuple[bytes, str]:
    wb, ws = _section_sheet("TuraVector - Recent Documents Report", "Recent Documents", "G", period_name,
                            ["Document Name", "Document Number", "Owner", "Status", "Created Date", "Document Type", "Brand"])
    status_map = {0: "Draft", 1: "Under Review", 2: "Approved", 3: "Revised"}
    for row, d in enumerate(documents, 5):
        ws.cell(row=row, column=1, value=d.document_name)
        ws.cell(row=row, column=2, value=d.document_number)
        ws.cell(row=row, column=3, value=_name(d.user))
        ws.cell(row=row, column=4, value=status_map.get(d.stage_status, "Unknown"))
        ws.cell(row=row, column=5, value=d.created_at.strftime("%Y-%m-%d %H:%M"))
        ws.cell(row=row, column=6, value=d.document_type.name if d.document_type else "N/A")
        ws.cell(row=row, column=7, value=d.brand.name if d.brand else "N/A")
    _autowidth(ws)
    return _save(wb), f"recent_documents_{_slug(period_name)}.xlsx"


def recent_documents_pdf(documents, period_name: str) -> tuple[bytes, str]:
    data = [["Document Name", "Owner", "Status", "Created", "Type"]]
    status_map = {0: "Draft", 1: "Review", 2: "Approved", 3: "Revised"}
    for d in documents[:20]:
        dt = d.document_type
        data.append([
            d.document_name[:25] + "..." if len(d.document_name) > 25 else d.document_name,
            _name(d.user),
            status_map.get(d.stage_status, "Unknown"),
            d.created_at.strftime("%Y-%m-%d"),
            dt.name[:10] + "..." if dt and len(dt.name) > 10 else (dt.name if dt else "N/A"),
        ])
    return (_section_pdf("TuraVector - Recent Documents Report", period_name, data, [2, 1.5, 1, 1, 1]),
            f"recent_documents_{_slug(period_name)}.pdf")


def top_users_excel(users: list[tuple], period_name: str) -> tuple[bytes, str]:
    wb, ws = _section_sheet("TuraVector - Top Users by Documents Report", "Top Users", "F", period_name,
                            ["Name", "Email", "User Type", "Document Count", "Status", "Joined Date"])
    for row, (u, count) in enumerate(users, 5):
        ws.cell(row=row, column=1, value=_name(u))
        ws.cell(row=row, column=2, value=u.email)
        ws.cell(row=row, column=3, value=user_type_display(u))
        ws.cell(row=row, column=4, value=count)
        ws.cell(row=row, column=5, value="Active" if u.is_active else "Inactive")
        ws.cell(row=row, column=6, value=u.date_joined.strftime("%Y-%m-%d"))
    _autowidth(ws)
    return _save(wb), f"top_users_{_slug(period_name)}.xlsx"


def top_users_pdf(users: list[tuple], period_name: str) -> tuple[bytes, str]:
    data = [["Name", "Email", "User Type", "Documents", "Status"]]
    for u, count in users[:15]:
        data.append([
            _name(u),
            u.email[:20] + "..." if len(u.email) > 20 else u.email,
            user_type_display(u),
            str(count),
            "Active" if u.is_active else "Inactive",
        ])
    return (_section_pdf("TuraVector - Top Users by Documents Report", period_name, data, [1.5, 2, 1, 0.8, 0.8]),
            f"top_users_{_slug(period_name)}.pdf")


def recent_users_excel(users, doc_counts: dict[int, int], period_name: str) -> tuple[bytes, str]:
    wb, ws = _section_sheet("TuraVector - Recent Users Report", "Recent Users", "G", period_name,
                            ["Name", "Email", "User Type", "Status", "Joined Date", "Last Login", "Document Count"])
    for row, u in enumerate(users, 5):
        ws.cell(row=row, column=1, value=_name(u))
        ws.cell(row=row, column=2, value=u.email)
        ws.cell(row=row, column=3, value=user_type_display(u))
        ws.cell(row=row, column=4, value="Active" if u.is_active else "Inactive")
        ws.cell(row=row, column=5, value=u.date_joined.strftime("%Y-%m-%d"))
        ws.cell(row=row, column=6, value=u.last_login.strftime("%Y-%m-%d") if u.last_login else "Never")
        ws.cell(row=row, column=7, value=doc_counts.get(u.id, 0))
    _autowidth(ws)
    return _save(wb), f"recent_users_{_slug(period_name)}.xlsx"


def recent_users_pdf(users, period_name: str) -> tuple[bytes, str]:
    data = [["Name", "Email", "User Type", "Status", "Joined"]]
    for u in users[:15]:
        data.append([
            _name(u),
            u.email[:20] + "..." if len(u.email) > 20 else u.email,
            user_type_display(u),
            "Active" if u.is_active else "Inactive",
            u.date_joined.strftime("%Y-%m-%d"),
        ])
    return (_section_pdf("TuraVector - Recent Users Report", period_name, data, [1.5, 2, 1, 1, 1]),
            f"recent_users_{_slug(period_name)}.pdf")
