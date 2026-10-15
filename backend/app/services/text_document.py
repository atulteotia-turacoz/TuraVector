"""Turn text written in the upload page's editor (HTML) into a .docx upload.

Keeps what the editor can produce: headings, paragraphs and their alignment / indent, bold, italic,
underline, strikethrough, sub/superscript, font, size, text colour and highlight, links, bulleted and
numbered lists, quotes, code blocks, horizontal lines, tables and pasted images (data: URLs).
"""
import base64
import io
import re
from html.parser import HTMLParser

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fastapi import UploadFile

BLOCKS = {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "blockquote", "pre"}
ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
         "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
# <font size="1..7"> and CSS keywords → points
FONT_SIZES = {"1": 8, "2": 10, "3": 12, "4": 14, "5": 18, "6": 24, "7": 36, "x-small": 8, "small": 10, "medium": 12,
              "large": 14, "x-large": 18, "xx-large": 24, "xxx-large": 36}
NAMED = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF", "yellow": "FFFF00",
         "orange": "FFA500", "purple": "800080", "gray": "808080", "grey": "808080"}
HIGHLIGHTS = {"FFFF00": WD_COLOR_INDEX.YELLOW, "00FF00": WD_COLOR_INDEX.BRIGHT_GREEN, "00FFFF": WD_COLOR_INDEX.TURQUOISE,
              "FF00FF": WD_COLOR_INDEX.PINK, "0000FF": WD_COLOR_INDEX.BLUE, "FF0000": WD_COLOR_INDEX.RED,
              "C0C0C0": WD_COLOR_INDEX.GRAY_25, "808080": WD_COLOR_INDEX.GRAY_50, "008000": WD_COLOR_INDEX.GREEN}


def _hex(value: str | None) -> str | None:
    if not value:
        return None
    v = value.strip().lower()
    if v in NAMED:
        return NAMED[v]
    m = re.match(r"#([0-9a-f]{3}|[0-9a-f]{6})$", v)
    if m:
        h = m.group(1)
        return (h if len(h) == 6 else "".join(c * 2 for c in h)).upper()
    m = re.match(r"rgba?\((\d+),\s*(\d+),\s*(\d+)", v)
    if m:
        return "".join(f"{int(x):02X}" for x in m.groups())
    return None


def _style(attrs: dict) -> dict:
    out = {}
    for part in (attrs.get("style") or "").split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            out[k.strip().lower()] = v.strip()
    return out


def _nearest_highlight(hexv: str):
    r, g, b = (int(hexv[i:i + 2], 16) for i in (0, 2, 4))
    best, dist = WD_COLOR_INDEX.YELLOW, 1e9
    for h, idx in HIGHLIGHTS.items():
        d = sum((a - int(h[i:i + 2], 16)) ** 2 for a, i in zip((r, g, b), (0, 2, 4)))
        if d < dist:
            best, dist = idx, d
    return best


def _size_pt(v: str | None) -> float | None:
    if not v:
        return None
    v = v.strip().lower()
    if v in FONT_SIZES:
        return FONT_SIZES[v]
    m = re.match(r"([\d.]+)(px|pt|em|rem)?$", v)
    if not m:
        return None
    n, unit = float(m.group(1)), m.group(2) or "px"
    return {"px": n * 0.75, "pt": n, "em": n * 12, "rem": n * 12}[unit]


def _add_hyperlink(paragraph, url: str, text: str, fmt: dict):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), fmt.get("color") or "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    if fmt.get("bold"):
        rpr.append(OxmlElement("w:b"))
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _horizontal_rule(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.append(bdr)


class _DocxBuilder(HTMLParser):
    def __init__(self, document):
        super().__init__(convert_charrefs=True)
        self.doc = document
        self.container = document           # document or a table cell
        self.para = None
        self.fmt_stack: list[tuple[str, dict]] = []
        self.lists: list[str] = []
        self.link: str | None = None
        self.block_style: list[dict] = []
        self.pre = 0
        # tables
        self.tables: list[dict] = []

    # ---------------------------------------------------------------- formatting
    def _fmt(self) -> dict:
        f: dict = {}
        for _, d in self.fmt_stack:
            f.update(d)
        return f

    def _inline(self, tag: str, attrs: dict) -> dict:
        st = _style(attrs)
        d: dict = {}
        if tag in ("b", "strong"):
            d["bold"] = True
        if tag in ("i", "em"):
            d["italic"] = True
        if tag == "u":
            d["underline"] = True
        if tag in ("s", "strike", "del"):
            d["strike"] = True
        if tag == "sub":
            d["sub"] = True
        if tag == "sup":
            d["sup"] = True
        if tag == "code":
            d["font"] = "Consolas"
        if tag == "font":
            if attrs.get("color"):
                d["color"] = _hex(attrs["color"])
            if attrs.get("face"):
                d["font"] = attrs["face"].split(",")[0].strip("'\" ")
            if attrs.get("size"):
                d["size"] = FONT_SIZES.get(attrs["size"])
        fw = st.get("font-weight", "")
        if fw in ("bold", "bolder") or (fw.isdigit() and int(fw) >= 600):
            d["bold"] = True
        if st.get("font-style") == "italic":
            d["italic"] = True
        deco = st.get("text-decoration", "") + " " + st.get("text-decoration-line", "")
        if "underline" in deco:
            d["underline"] = True
        if "line-through" in deco:
            d["strike"] = True
        if st.get("vertical-align") == "sub":
            d["sub"] = True
        if st.get("vertical-align") == "super":
            d["sup"] = True
        if _hex(st.get("color")):
            d["color"] = _hex(st.get("color"))
        bg = _hex(st.get("background-color"))
        if bg and bg != "FFFFFF":
            d["highlight"] = bg
        if st.get("font-family"):
            d["font"] = st["font-family"].split(",")[0].strip("'\" ")
        if _size_pt(st.get("font-size")):
            d["size"] = _size_pt(st.get("font-size"))
        return {k: v for k, v in d.items() if v}

    # ---------------------------------------------------------------- paragraphs
    def _new(self, tag: str = "p", attrs: dict | None = None):
        attrs = attrs or {}
        target = self.container
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6") and target is self.doc:
            self.para = self.doc.add_heading(level=int(tag[1]))
        elif tag == "li":
            ordered = self.lists[-1:] == ["ol"]
            depth = min(len(self.lists), 3)
            style = ("List Number" if ordered else "List Bullet") + ("" if depth <= 1 else f" {depth}")
            try:
                self.para = target.add_paragraph(style=style)
            except KeyError:
                self.para = target.add_paragraph(style="List Number" if ordered else "List Bullet")
        elif tag == "blockquote":
            self.para = target.add_paragraph(style="Quote") if target is self.doc else target.add_paragraph()
        else:
            self.para = target.add_paragraph()
        st = _style(attrs)
        align = (st.get("text-align") or attrs.get("align") or "").lower()
        for bs in self.block_style:
            align = align or bs.get("align", "")
        if align in ALIGN:
            self.para.alignment = ALIGN[align]
        indent = st.get("margin-left") or st.get("padding-left")
        if indent and _size_pt(indent):
            self.para.paragraph_format.left_indent = Pt(_size_pt(indent))
        if tag == "pre":
            self.fmt_stack.append(("pre", {"font": "Consolas", "size": 10}))

    def _ensure(self):
        if self.para is None:
            self._new()

    # ---------------------------------------------------------------- parser
    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in ("ul", "ol"):
            self.lists.append(tag)
        elif tag == "table":
            self.tables.append({"rows": [], "row": None, "cell": None, "outer": (self.container, self.para)})
        elif tag == "tr" and self.tables:
            self.tables[-1]["rows"].append([])
        elif tag in ("td", "th") and self.tables:
            t = self.tables[-1]
            if not t["rows"]:
                t["rows"].append([])
            t["rows"][-1].append({"html": [], "header": tag == "th"})
        elif tag == "hr":
            p = self.container.add_paragraph()
            _horizontal_rule(p)
            self.para = None
        elif tag == "img":
            self._image(attrs.get("src") or "", attrs)
        elif tag in BLOCKS:
            if self.tables:
                self._cell_text("\n")
                return
            self.block_style.append({"align": (_style(attrs).get("text-align") or attrs.get("align") or "").lower()})
            self._new(tag, attrs)
            if tag == "pre":
                self.pre += 1
        elif tag == "br":
            if self.tables:
                self._cell_text("\n")
                return
            self._ensure()
            self.para.add_run().add_break()
        elif tag == "a":
            self.link = attrs.get("href") or None
        else:
            d = self._inline(tag, attrs)
            self.fmt_stack.append((tag, d))

    def handle_endtag(self, tag):
        if tag in ("ul", "ol") and self.lists:
            self.lists.pop()
            self.para = None
        elif tag == "table" and self.tables:
            self._flush_table(self.tables.pop())
        elif tag in BLOCKS:
            if self.tables:
                return
            if tag == "pre":
                self.pre = max(self.pre - 1, 0)
                self._pop("pre")
            if self.block_style:
                self.block_style.pop()
            self.para = None
        elif tag == "a":
            self.link = None
        else:
            self._pop(tag)

    def _pop(self, tag):
        for i in range(len(self.fmt_stack) - 1, -1, -1):
            if self.fmt_stack[i][0] == tag:
                del self.fmt_stack[i]
                return

    def handle_data(self, data):
        if self.tables:
            self._cell_text(data)
            return
        if not self.pre:
            data = re.sub(r"\s+", " ", data)
        if not data.strip() and self.para is None:
            return
        self._ensure()
        f = self._fmt()
        if self.link and re.match(r"https?://|mailto:", self.link):
            _add_hyperlink(self.para, self.link, data, f)
            return
        run = self.para.add_run(data)
        run.bold = f.get("bold") or None
        run.italic = f.get("italic") or None
        run.underline = f.get("underline") or None
        if f.get("strike"):
            run.font.strike = True
        if f.get("sub"):
            run.font.subscript = True
        if f.get("sup"):
            run.font.superscript = True
        if f.get("color"):
            run.font.color.rgb = RGBColor.from_string(f["color"])
        if f.get("highlight"):
            run.font.highlight_color = _nearest_highlight(f["highlight"])
        if f.get("size"):
            run.font.size = Pt(f["size"])
        if f.get("font"):
            run.font.name = f["font"]

    # ---------------------------------------------------------------- tables & images
    def _cell_text(self, text: str):
        t = self.tables[-1]
        if t["rows"] and t["rows"][-1]:
            t["rows"][-1][-1]["html"].append(text)

    def _flush_table(self, t: dict):
        rows = [r for r in t["rows"] if r]
        if not rows:
            return
        cols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                text = re.sub(r"[ \t]+", " ", "".join(c["html"])).strip()
                cell = table.cell(i, j)
                cell.text = text
                if c["header"] and cell.paragraphs and cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].bold = True
        self.para = None

    def _image(self, src: str, attrs: dict):
        m = re.match(r"data:image/(png|jpe?g|gif|bmp);base64,(.+)$", src, re.S)
        if not m:
            return
        try:
            data = base64.b64decode(m.group(2))
        except ValueError:
            return
        width = _size_pt(_style(attrs).get("width") or attrs.get("width"))
        inches = min((width or 432) / 72, 6.0)  # 432pt = 6in, the usable page width
        try:
            p = self.container.add_paragraph()
            p.add_run().add_picture(io.BytesIO(data), width=Inches(inches))
        except Exception:  # noqa: BLE001 — an unreadable image is skipped, the text is kept
            return
        self.para = None


def plain_text(html: str) -> str:
    text = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html or "")).strip()
    return text or ("image" if "<img" in (html or "") else "")


def html_to_docx(html: str, title: str) -> UploadFile:
    from docx import Document as Docx
    doc = Docx()
    doc.core_properties.title = title
    _DocxBuilder(doc).feed(html or "")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    safe = re.sub(r"[^A-Za-z0-9 _.-]+", "", title).strip()[:80] or "document"
    return UploadFile(file=buf, filename=f"{safe}.docx")
