"""Review work of reviewers, by university / organisation — for the superadmin's team pages and the
downloadable report.

A review assignment belongs to the workspace of the document under review (the document owner's
workspace; documents uploaded by platform superadmins count as the home workspace's, SMBT).
"""
from __future__ import annotations

import csv
import io
from collections import defaultdict

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.tenancy import HOME_WORKSPACE_SLUG, unfiltered
from app.core.utils import iso
from app.models import AssignedReviewer, Document, DocumentVersion, Organisation, OrgMembership, User, utcnow

STATUSES = ("completed", "overdue", "in_progress", "pending")
STATUS_LABELS = {"completed": "Completed", "overdue": "Overdue", "in_progress": "In progress", "pending": "Pending"}
EMPTY = {"assigned": 0, "completed": 0, "pending": 0, "in_progress": 0, "overdue": 0}


def status_of(a: AssignedReviewer, now) -> str:
    if a.is_completed:
        return "completed"
    if a.review_date is not None and a.review_date < now:
        return "overdue"
    return "in_progress" if a.started_at is not None else "pending"


def assignments(db: Session, user_ids: list[int] | None, organisation_id: int | None = None) -> list[dict]:
    """Every active review assignment of these reviewers (or of everyone, with `organisation_id`), with
    the document and its workspace. `organisation_id` keeps only reviews of that workspace's documents."""
    if user_ids is not None and not user_ids:
        return []
    now = utcnow()
    with unfiltered(db):
        home = db.execute(select(Organisation).where(Organisation.slug == HOME_WORKSPACE_SLUG)).scalar_one_or_none()
        rows = db.execute(
            select(AssignedReviewer, DocumentVersion, Document, Organisation, User.is_superuser)
            .join(DocumentVersion, DocumentVersion.id == AssignedReviewer.document_version_id)
            .join(Document, Document.id == DocumentVersion.document_id)
            .join(User, User.id == Document.user_id)
            .outerjoin(OrgMembership, OrgMembership.user_id == Document.user_id)
            .outerjoin(Organisation, Organisation.id == OrgMembership.organisation_id)
            .where(AssignedReviewer.is_active == 1,
                   *([AssignedReviewer.user_id.in_(user_ids)] if user_ids is not None else []))
            .order_by(AssignedReviewer.review_date.desc())
        ).all()
    out = []
    for a, v, d, org, owner_is_superuser in rows:
        if org is None and owner_is_superuser and home is not None:
            org = home
        if organisation_id is not None and (org is None or org.id != organisation_id):
            continue
        status = status_of(a, now)
        out.append({
            "id": a.id, "reviewer_id": a.user_id,
            "organisation_id": org.id if org else None, "organisation": org.name if org else "No workspace",
            "document_id": d.id, "document_number": d.document_number, "document_name": d.document_name,
            "version": v.version_number, "assigned_at": iso(a.created_at), "review_date": iso(a.review_date),
            "started_at": iso(a.started_at), "completed_at": iso(a.completed_at),
            "status": status, "status_label": STATUS_LABELS[status], "remarks": a.remarks or "",
        })
    return out


def tally(items: list[dict]) -> dict:
    t = dict(EMPTY)
    for i in items:
        t["assigned"] += 1
        if i["status"] == "completed":
            t["completed"] += 1
        else:
            t["pending"] += 1
            t["in_progress"] += i["started_at"] is not None  # started (late or not)
            t["overdue"] += i["status"] == "overdue"
    return t


def by_organisation(items: list[dict]) -> list[dict]:
    groups: dict[tuple, list[dict]] = defaultdict(list)
    for i in items:
        groups[(i["organisation_id"], i["organisation"])].append(i)
    return sorted(({"organisation_id": oid, "organisation": name, **tally(g)} for (oid, name), g in groups.items()),
                  key=lambda r: (-r["assigned"], r["organisation"]))


def per_reviewer(db: Session, user_ids: list[int]) -> dict[int, dict]:
    """{reviewer id: {"totals": {...}, "by_organisation": [...]}}"""
    items = assignments(db, user_ids)
    grouped: dict[int, list[dict]] = defaultdict(list)
    for i in items:
        grouped[i["reviewer_id"]].append(i)
    return {uid: {"totals": tally(g), "by_organisation": by_organisation(g)} for uid, g in grouped.items()}


# ---------------------------------------------------------------------------------------------- report files
def _people(db: Session, user_ids: list[int]) -> dict[int, User]:
    with unfiltered(db):
        return {u.id: u for u in db.execute(select(User).where(User.id.in_(user_ids))).scalars()}


def _name(u: User) -> str:
    return f"{u.first_name} {u.last_name}".strip() or u.username


def report_tables(db: Session, user_ids: list[int], assigned_to: dict[int, list[str]],
                  items: list[dict] | None = None) -> dict[str, list[list]]:
    """The three report tables: per reviewer, per reviewer × university, and every assignment.
    `items` limits the report to those assignments (e.g. one workspace's documents)."""
    people = _people(db, user_ids)
    items = assignments(db, user_ids) if items is None else items
    grouped: dict[int, list[dict]] = defaultdict(list)
    for i in items:
        grouped[i["reviewer_id"]].append(i)
    order = sorted(user_ids, key=lambda uid: _name(people[uid]).lower() if uid in people else "")
    counts = ["Assigned", "Completed", "Pending", "Started", "Overdue"]

    summary = [["Reviewer", "Username", "Email", "Assigned to", *counts]]
    by_org = [["Reviewer", "Username", "University / organisation", *counts]]
    for uid in order:
        u = people.get(uid)
        if u is None:
            continue
        t = tally(grouped.get(uid, []))
        summary.append([_name(u), u.username, u.email, ", ".join(assigned_to.get(uid, [])),
                        t["assigned"], t["completed"], t["pending"], t["in_progress"], t["overdue"]])
        for r in by_organisation(grouped.get(uid, [])):
            by_org.append([_name(u), u.username, r["organisation"],
                           r["assigned"], r["completed"], r["pending"], r["in_progress"], r["overdue"]])

    detail = [["Reviewer", "Username", "University / organisation", "Document number", "Document", "Version",
               "Status", "Assigned on", "Review date", "Started", "Completed", "Remarks"]]
    for i in sorted(items, key=lambda x: (_name(people[x["reviewer_id"]]).lower(), x["organisation"], x["review_date"] or "")):
        u = people[i["reviewer_id"]]
        detail.append([_name(u), u.username, i["organisation"], i["document_number"], i["document_name"], i["version"],
                       i["status_label"], _day(i["assigned_at"]), _day(i["review_date"]), _day(i["started_at"]),
                       _day(i["completed_at"]), i["remarks"]])
    return {"Summary": summary, "By university": by_org, "Assignments": detail}


def _day(value: str | None) -> str:
    return value[:10] if value else ""


def as_xlsx(tables: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in tables.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def as_csv(rows: list[list]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8 names correctly


def workspace_reviewers(db: Session, organisation_id: int) -> tuple[list[dict], dict[int, User]]:
    """Reviews of one workspace's documents, and everyone who reviews them (members or Turacoz staff)."""
    items = assignments(db, None, organisation_id)
    return items, _people(db, sorted({i["reviewer_id"] for i in items}))
