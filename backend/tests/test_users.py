"""Users module tests. All writes happen inside the rolled-back test transaction (see conftest.py)."""
import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.security import check_password
from app.models import (
    Organisation,
    EmailLog, Permission, Role, RolePermission, TemporaryPassword, User, UserPermission, UserRole, UserType,
)
from app.services.users_common import create_user


@pytest.fixture()
def admin(login_as):
    return login_as("admin")


@pytest.fixture()
def student_type(db):
    return db.execute(select(UserType).where(UserType.code == "student")).scalar_one()


def _new_user(db, username="tmp_users_test", **kw):
    u = create_user(db, username=username, email=f"{username}@example.com", password="Secret#123",
                    first_name="Tmp", last_name="User", **kw)
    db.commit()
    return u


# ------------------------------------------------------------------ access control
def test_non_manager_is_refused(client, login_as):
    h = login_as("reviewer1")
    r = client.get("/api/users/manage", headers=h)
    assert r.status_code == 403
    assert r.json()["detail"]["message"] == "You do not have permission to manage users"
    assert client.get("/api/users/reports", headers=h).status_code == 403
    assert client.post("/api/users/create", json={}, headers=h).status_code == 403


def test_dashboard_and_lists(client, admin):
    d = client.get("/api/users/manager-dashboard", headers=admin).json()
    assert {"total_documents", "draft_documents", "pending_reviews", "recent_documents"} <= d.keys()
    assert client.get("/api/users/manager-dashboard?search=draft", headers=admin).status_code == 200
    assert client.get("/api/users/manager-dashboard?search=student&page=999", headers=admin).status_code == 200
    m = client.get("/api/users/manage?status=active", headers=admin).json()
    assert m["users"] and all(u["is_active"] for u in m["users"])
    assert [s["code"] for s in m["user_type_stats"]] == ["student", "teacher", "reviewer", "manager"]


# ------------------------------------------------------------------ create / edit / toggle / delete
def test_create_user_validation(client, admin):
    r = client.post("/api/users/create", json={"username": "admin", "password": "short", "confirm_password": "x"}, headers=admin)
    assert r.status_code == 400
    errs = r.json()["detail"]
    assert "Username already exists" in errs and "Email is required" in errs
    assert "Password must be at least 8 characters long" in errs and "Passwords do not match" in errs
    assert "User type is required" in errs


def _smbt(db) -> int:
    return db.execute(select(Organisation.id).where(Organisation.slug == "smbt")).scalar_one()


def test_create_user(client, admin, db, student_type, sent_emails):
    role = db.execute(select(Role).where(Role.is_active == 1)).scalars().first()
    body = {"username": "new_mgr_user", "email": "new_mgr_user@Example.COM", "first_name": "New", "last_name": "Person",
            "user_type": str(student_type.id), "password": "Password1!", "confirm_password": "Password1!",
            "is_manager": True, "role": str(role.id), "phone": "", "address": "Somewhere", "organisation": str(_smbt(db))}
    r = client.post("/api/users/create", json=body, headers=admin)
    assert r.status_code == 200, r.text
    assert r.json()["messages"][0]["text"] == 'User "New Person" created successfully!'
    u = db.execute(select(User).where(User.username == "new_mgr_user")).scalar_one()
    assert check_password("Password1!", u.password)
    assert u.email == "new_mgr_user@example.com" and u.is_manager == 1 and u.phone is None
    assert u.user_type == "student" and u.user_type_new_id == student_type.id
    assert db.execute(select(TemporaryPassword).where(TemporaryPassword.user_id == u.id)).scalar_one().password == "Password1!"
    assert db.execute(select(UserRole).where(UserRole.user_id == u.id, UserRole.role_id == role.id)).scalar_one()
    assert sent_emails and sent_emails[-1]["subject"] == "Welcome to TuraVector - Your Account Has Been Created"


def test_edit_toggle_delete_user(client, admin, db, student_type):
    u = _new_user(db)
    r = client.get(f"/api/users/{u.id}/edit", headers=admin)
    assert r.status_code == 200
    r = client.post(f"/api/users/{u.id}/edit", headers=admin, json={
        "first_name": "Changed", "last_name": "Name", "email": u.email, "username": u.username, "phone": "123",
        "address": "", "user_type": str(student_type.id), "is_manager": False, "is_active": True})
    assert r.status_code == 200 and r.json()["message"] == "User Changed Name updated successfully"
    db.refresh(u)
    assert u.first_name == "Changed" and u.phone == "123"

    r = client.post(f"/api/users/{u.id}/edit", headers=admin, json={"user_type": "999999", "is_active": True})
    assert r.status_code == 400 and r.json()["detail"] == "Invalid user type selected"

    r = client.post(f"/api/users/{u.id}/toggle-status", headers=admin)
    assert r.json()["message"] == "User Changed Name has been deactivated"
    r = client.post(f"/api/users/{u.id}/toggle-status", headers=admin)
    assert r.json()["message"] == "User Changed Name has been activated"
    admin_id = db.execute(select(User.id).where(User.username == "admin")).scalar_one()
    r = client.post(f"/api/users/{admin_id}/toggle-status", headers=admin)
    assert r.status_code == 400 and r.json()["detail"] == "You cannot deactivate your own account"

    r = client.post(f"/api/users/{u.id}/delete", headers=admin)
    assert r.json()["message"] == "User Changed Name has been deactivated"
    db.refresh(u)
    assert u.is_active == 0
    assert client.get(f"/api/users/{u.id}/profile", headers=admin).json()["viewed_user"]["is_active"] is False


# ------------------------------------------------------------------ passwords
def test_change_and_view_password(client, admin, db, sent_emails):
    u = _new_user(db)
    r = client.post(f"/api/users/{u.id}/change-password", headers=admin,
                    json={"new_password": "abc", "confirm_password": "abd"})
    assert r.status_code == 400
    assert r.json()["detail"] == ["Password must be at least 8 characters long", "Passwords do not match"]

    r = client.post(f"/api/users/{u.id}/change-password", headers=admin,
                    json={"new_password": "NewPass#2024", "confirm_password": "NewPass#2024", "send_email": True, "force_change": True})
    assert r.status_code == 200
    assert r.json()["message"] == "Password changed successfully for Tmp User and notification email sent"
    db.refresh(u)
    assert check_password("NewPass#2024", u.password)
    log = db.execute(select(EmailLog).where(EmailLog.recipient_user_id == u.id)).scalars().all()[-1]
    assert log.email_type == "password_reset" and log.status == "sent"
    assert "New Password: NewPass#2024" in sent_emails[-1]["message"]

    v = client.get(f"/api/users/{u.id}/view-password", headers=admin).json()
    assert v["password_info"]["password"] == "NewPass#2024"


def test_view_password_superuser_only(client, login_as, db):
    u = _new_user(db, is_manager=True)
    r = client.get(f"/api/users/{u.id}/view-password", headers=login_as(u.username))
    assert r.status_code == 403 and r.json()["detail"]["message"] == "Only administrators can view user passwords"


# ------------------------------------------------------------------ bulk import
def _csv(rows: list[str]) -> bytes:
    return ("\n".join(["username,email,first_name,last_name,user_type,password,phone,is_manager"] + rows) + "\n").encode()


def test_bulk_import_csv(client, admin, db, sent_emails):
    data = _csv(["bulk_a,bulk_a@example.com,Bulk,A,Student,Passw0rd!,,true",
                 "bulk_b,bulk_b@example.com,Bulk,B,teacher,,555,false",
                 "admin,dupe@example.com,Dup,User,Student,,,"])
    r = client.post("/api/users/bulk-import", headers=admin,
                    files={"file": ("users.csv", data, "text/csv")},
                    data={"send_welcome_email": "true", "skip_errors": "true", "organisation": str(_smbt(db))})
    out = r.json()
    assert out["success"] and out["redirect"] == "manage_users"
    texts = [m["text"] for m in out["messages"]]
    assert "Successfully imported 2 users." in texts and "Failed to import 1 users." in texts
    assert 'Row 4: Username "admin" already exists' in texts
    a = db.execute(select(User).where(User.username == "bulk_a")).scalar_one()
    assert check_password("Passw0rd!", a.password) and a.is_manager == 1 and a.user_type == "student"
    b = db.execute(select(User).where(User.username == "bulk_b")).scalar_one()
    assert b.user_type == "teacher" and b.phone == "555.0"  # str(float) — same as Django + pandas
    assert len([e for e in sent_emails if e["subject"] == "Welcome to TuraVector"]) == 2


def test_bulk_import_stops_on_error_and_rolls_back(client, admin, db):
    data = _csv(["bulk_c,bulk_c@example.com,Bulk,C,Student,,,", "bulk_d,bulk_d@example.com,Bulk,D,NoSuchType,,,"])
    r = client.post("/api/users/bulk-import", headers=admin, files={"file": ("u.csv", data, "text/csv")},
                    data={"send_welcome_email": "false", "skip_errors": "false", "organisation": str(_smbt(db))})
    out = r.json()
    assert not out["success"] and out["redirect"] == "bulk_import_users"
    assert out["messages"][0]["text"] == 'Error processing file: Row 3: Invalid user type "NoSuchType"'
    assert db.execute(select(User).where(User.username == "bulk_c")).scalar_one_or_none() is None


def test_bulk_import_xlsx_and_validation(client, admin, db):
    wb = Workbook()
    ws = wb.active
    ws.append(["username", "email", "first_name", "last_name", "user_type"])
    ws.append(["bulk_x", "bulk_x@example.com", "Bulk", "X", "Reviewer"])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/users/bulk-import", headers=admin,
                    files={"file": ("u.xlsx", buf.getvalue(), "application/octet-stream")},
                    data={"organisation": str(_smbt(db))})
    assert r.json()["messages"][0]["text"] == "Successfully imported 1 users."
    assert db.execute(select(User).where(User.username == "bulk_x")).scalar_one().user_type == "reviewer"

    r = client.post("/api/users/bulk-import", headers=admin, files={"file": ("u.txt", b"x", "text/plain")})
    assert r.status_code == 400 and "not allowed" in r.json()["detail"]["message"]
    r = client.post("/api/users/bulk-import", headers=admin, files={"file": ("u.csv", b"username,email\na,b\n", "text/csv")},
                    data={"organisation": str(_smbt(db))})
    assert r.json()["messages"][0]["text"] == "Missing required columns: first_name, last_name, user_type"


def test_download_template(client, admin):
    r = client.get("/api/users/bulk-import/template?format=csv", headers=admin)
    assert r.status_code == 200 and r.text.startswith("username,email,first_name")
    assert "REQUIRED: Unique username" in r.text
    r = client.get("/api/users/bulk-import/template?format=xlsx", headers=admin)
    ws = load_workbook(io.BytesIO(r.content))["Users"]
    assert ws["A1"].value == "username" and ws["A2"].value == "REQUIRED: Unique username"


# ------------------------------------------------------------------ roles & permissions (legacy)
def test_legacy_role_crud_and_user_permissions(client, admin, db):
    perms = db.execute(select(Permission).where(Permission.is_active == 1).limit(3)).scalars().all()
    r = client.post("/api/rbac/legacy/roles/create", headers=admin,
                    json={"name": "Zz Test Role", "description": "d", "permissions": [str(p.id) for p in perms[:2]]})
    assert r.status_code == 200 and r.json()["message"] == 'Role "Zz Test Role" has been created successfully'
    role_id = r.json()["id"]
    assert len(db.execute(select(RolePermission).where(RolePermission.role_id == role_id)).scalars().all()) == 2
    r = client.post("/api/rbac/legacy/roles/create", headers=admin, json={"name": "Zz Test Role"})
    assert r.json()["detail"] == 'Role "Zz Test Role" already exists. Please choose a different name.'

    form = client.get(f"/api/rbac/legacy/roles/{role_id}", headers=admin).json()
    assert sorted(form["current_permission_ids"]) == sorted(p.id for p in perms[:2])
    r = client.post(f"/api/rbac/legacy/roles/{role_id}", headers=admin,
                    json={"name": "Zz Test Role 2", "description": "", "permissions": [str(perms[2].id), "999999"]})
    assert r.json()["message"] == 'Role "Zz Test Role 2" has been updated successfully'
    assert [rp.permission_id for rp in db.execute(select(RolePermission).where(RolePermission.role_id == role_id)).scalars()] == [perms[2].id]

    u = _new_user(db)
    avail = client.get(f"/api/rbac/legacy/roles/{role_id}/available-users", headers=admin).json()["users"]
    assert any(x["id"] == u.id for x in avail)
    r = client.post(f"/api/rbac/legacy/roles/{role_id}/assign-user", headers=admin,
                    json={"user_id": str(u.id), "expires_at": "2099-01-01T10:00"})
    assert r.json()["message"] == 'Role "Zz Test Role 2" has been assigned to Tmp User'
    r = client.post(f"/api/rbac/legacy/roles/{role_id}/assign-user", headers=admin, json={"user_id": str(u.id)})
    assert r.json()["messages"][0]["level"] == "warning"

    # per-user overrides
    p = perms[0]
    r = client.post(f"/api/users/{u.id}/permissions", headers=admin,
                    json={"action": "deny_permission", "permission_id": str(p.id), "reason": "test"})
    assert r.json()["message"].startswith(f'Permission "') and "denied for Tmp User" in r.json()["message"]
    # (new users also get an automatic Content Review card row, so look up the permission set here)
    up = db.execute(select(UserPermission).where(UserPermission.user_id == u.id, UserPermission.permission_id == p.id)).scalar_one()
    assert up.is_granted == 0 and up.reason == "test"
    client.post(f"/api/users/{u.id}/permissions", headers=admin, json={"action": "grant_permission", "permission_id": str(p.id)})
    db.refresh(up)
    assert up.is_granted == 1
    client.post(f"/api/users/{u.id}/permissions", headers=admin, json={"action": "remove_permission", "permission_id": str(p.id)})
    assert db.execute(select(UserPermission).where(UserPermission.user_id == u.id, UserPermission.permission_id == p.id)).scalar_one_or_none() is None
    page = client.get(f"/api/users/{u.id}/permissions", headers=admin).json()
    assert page["user_roles"][0]["role"]["id"] == role_id
    client.post(f"/api/users/{u.id}/permissions", headers=admin, json={"action": "remove_role", "role_id": str(role_id)})
    assert db.execute(select(UserRole).where(UserRole.user_id == u.id)).scalar_one_or_none() is None

    r = client.post("/api/rbac/legacy/roles", headers=admin, json={"action": "delete_role", "role_id": str(role_id)})
    assert r.json()["message"] == 'Role "Zz Test Role 2" deleted successfully'
    assert db.get(Role, role_id) is None
    sys_role = db.execute(select(Role).where(Role.is_system_role == 1)).scalars().first()
    r = client.post("/api/rbac/legacy/roles", headers=admin, json={"action": "delete_role", "role_id": str(sys_role.id)})
    assert r.status_code == 400 and r.json()["detail"] == "System roles cannot be deleted"


# ------------------------------------------------------------------ comprehensive permission screens
def test_comprehensive_roles(client, admin, db):
    assert client.get("/api/rbac/dashboard", headers=admin).json()["stats"]["total_roles"] >= 1
    assert "roles" in client.get("/api/rbac/roles", headers=admin).json()
    r = client.post("/api/rbac/roles", headers=admin, json={"name": "  Zz Comp Role  ", "description": " x "})
    assert r.json()["message"] == 'Role "Zz Comp Role" created successfully.'
    role_id = r.json()["id"]
    perm = db.execute(select(Permission).where(Permission.is_active == 1)).scalars().first()

    r = client.post(f"/api/rbac/roles/{role_id}/assign-permission", headers=admin, json={"permission_id": perm.id})
    assert r.json()["message"] == f'Permission "{perm.name}" assigned to role "Zz Comp Role"'
    r = client.post(f"/api/rbac/roles/{role_id}/assign-permission", headers=admin, json={"permission_id": perm.id})
    assert r.status_code == 400 and r.json()["detail"] == "Permission already assigned to this role"
    detail = client.get(f"/api/rbac/roles/{role_id}", headers=admin).json()
    assert detail["permissions_by_dashboard"][0]["modules"][0]["permissions"][0]["id"] == perm.id
    r = client.delete(f"/api/rbac/roles/{role_id}/permissions/{perm.id}", headers=admin)
    assert r.json()["message"] == f'Permission "{perm.name}" removed from role "Zz Comp Role"'
    r = client.delete(f"/api/rbac/roles/{role_id}/permissions/{perm.id}", headers=admin)
    assert r.status_code == 400

    r = client.post(f"/api/rbac/roles/{role_id}/edit", headers=admin, json={"name": "Zz Comp Role B", "permissions": [str(perm.id)]})
    assert r.json()["message"] == 'Role "Zz Comp Role B" has been updated successfully'

    u = _new_user(db)
    r = client.post(f"/api/rbac/users/{u.id}/roles", headers=admin, json={"action": "assign", "role_id": str(role_id)})
    assert r.json()["message"] == 'Role "Zz Comp Role B" assigned to Tmp User'
    eff = client.get(f"/api/rbac/users/{u.id}/permissions", headers=admin).json()
    assert eff["total_permissions"] == 1
    r = client.post(f"/api/rbac/users/{u.id}/roles", headers=admin, json={"action": "remove", "role_id": str(role_id)})
    assert r.json()["message"] == 'Role "Zz Comp Role B" removed from Tmp User'


def test_dashboard_permission_toggle(client, admin, db):
    access = client.get("/api/rbac/dashboard-access", headers=admin).json()
    assert access["can_edit"] is True
    perm_id = access["content_lab_permission"]["id"]
    role = Role(name="Zz Toggle Role", description="", is_system_role=0, is_active=1)
    db.add(role)
    db.commit()
    body = {"role_id": role.id, "permission_id": perm_id, "action": "grant"}
    r = client.post("/api/rbac/toggle-role-dashboard-permission", headers=admin, json=body)
    assert r.json()["message"] == 'Dashboard access "Content Review - Dashboard Access" granted to role "Zz Toggle Role"'
    assert client.post("/api/rbac/toggle-role-dashboard-permission", headers=admin, json=body).json()["message"] == \
        'Role "Zz Toggle Role" already has this permission'
    r = client.post("/api/rbac/toggle-role-dashboard-permission", headers=admin, json={**body, "action": "revoke"})
    assert r.json()["message"] == 'Dashboard access "Content Review - Dashboard Access" revoked from role "Zz Toggle Role"'
    assert client.post("/api/rbac/toggle-role-dashboard-permission", headers=admin, json={**body, "action": "x"}).status_code == 400


# ------------------------------------------------------------------ user types
def test_user_types_crud(client, admin, db):
    lst = client.get("/api/users/user-types", headers=admin).json()["user_types"]
    assert any(t["code"] == "student" and t["user_count"] > 0 for t in lst)
    r = client.post("/api/users/user-types", headers=admin, json={"name": "Zz Type", "code": "Zz Type", "is_active": True})
    assert r.json()["message"] == 'User type "Zz Type" created successfully.'
    ut = db.get(UserType, r.json()["id"])
    assert ut.code == "zz_type" and ut.is_system == 0
    r = client.post("/api/users/user-types", headers=admin, json={"name": "zz type", "code": "other"})
    assert r.json()["detail"] == "A user type with this name already exists."
    r = client.post(f"/api/users/user-types/{ut.id}", headers=admin,
                    json={"name": "Zz Type 2", "code": "zz_type2", "description": "d", "is_active": False})
    assert r.json()["message"] == 'User type "Zz Type 2" updated successfully.'
    db.refresh(ut)
    assert ut.is_active == 0
    r = client.post(f"/api/users/user-types/{ut.id}/toggle-status", headers=admin)
    assert r.json() == {"success": True, "message": 'User type "Zz Type 2" activated successfully.', "is_active": True}

    u = _new_user(db, user_type_new_id=ut.id)
    r = client.post(f"/api/users/user-types/{ut.id}/delete", headers=admin)
    assert r.json()["detail"] == 'Cannot delete user type "Zz Type 2" as it is assigned to 1 user(s).'
    u.user_type_new_id = None
    db.commit()
    r = client.post(f"/api/users/user-types/{ut.id}/delete", headers=admin)
    assert r.json()["message"] == 'User type "Zz Type 2" deleted successfully.'
    student = db.execute(select(UserType).where(UserType.code == "student")).scalar_one()
    r = client.post(f"/api/users/user-types/{student.id}/delete", headers=admin)
    assert r.json()["detail"] == "System user types cannot be deleted."


# ------------------------------------------------------------------ reports & exports
def test_reports_and_exports(client, admin):
    rep = client.get("/api/users/reports?date_range=7", headers=admin).json()
    assert {"total_documents", "top_users", "recent_users", "recent_documents"} <= rep.keys()

    r = client.get("/api/users/reports/export?format=excel&date_range=30", headers=admin)
    assert r.status_code == 200 and 'turavector_report_last_30_days.xlsx' in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content))["TuraVector Report"]
    assert ws["A1"].value == "TuraVector Analytics Report" and ws["A4"].value == "Key Metrics"

    r = client.get("/api/users/reports/export?format=pdf&date_range=365", headers=admin)
    assert r.content.startswith(b"%PDF") and "turavector_report_last_year.pdf" in r.headers["content-disposition"]
    assert client.get("/api/users/reports/export?format=csv", headers=admin).status_code == 400

    for section, sheet in [("recent_documents", "Recent Documents"), ("top_users", "Top Users"), ("recent_users", "Recent Users")]:
        r = client.get(f"/api/users/reports/export-section?section={section}&format=excel&date_range=90", headers=admin)
        assert r.status_code == 200
        assert load_workbook(io.BytesIO(r.content))[sheet]["A4"].value is not None
        r = client.get(f"/api/users/reports/export-section?section={section}&format=pdf", headers=admin)
        assert r.content.startswith(b"%PDF")
    assert client.get("/api/users/reports/export-section?section=nope", headers=admin).status_code == 400


def test_reports_custom_date_range(client, admin, db):
    from app.models import Document
    first = db.execute(select(Document).where(Document.is_active == 1).order_by(Document.created_at)).scalars().first()
    day = first.created_at.date().isoformat()
    r = client.get("/api/users/reports", headers=admin,
                   params={"date_range": "custom", "date_from": "2000-01-01", "date_to": "2000-01-31"}).json()
    assert r["documents_in_period"] == 0 and r["recent_documents"] == [] and r["top_users"] == []
    assert r["period_name"] == "01 Jan 2000 to 31 Jan 2000"
    wide = client.get("/api/users/reports", headers=admin,
                      params={"date_range": "custom", "date_from": "2000-01-01", "date_to": "2100-01-01"}).json()
    assert wide["documents_in_period"] == wide["total_documents"]  # every document falls in the range
    assert wide["recent_documents"] and all(day <= d["created_at"][:10] for d in wide["recent_documents"])
    bad = client.get("/api/users/reports", headers=admin, params={"date_range": "custom", "date_from": "2024-05-01", "date_to": "2024-04-01"})
    assert bad.status_code == 400 and "on or after" in bad.json()["detail"]
    x = client.get("/api/users/reports/export", headers=admin,
                   params={"format": "excel", "date_range": "custom", "date_from": "2000-01-01", "date_to": "2000-01-31"})
    assert x.status_code == 200
    s = client.get("/api/users/reports/export-section", headers=admin,
                   params={"section": "top_users", "format": "excel", "date_range": "custom", "date_from": "2000-01-01", "date_to": "2000-01-31"})
    assert s.status_code == 200
