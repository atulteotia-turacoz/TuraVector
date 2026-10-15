"""Turacoz team (system users) assigned to workspaces, and workspace points of contact."""
import uuid

import pytest
from sqlalchemy import select

from app.core.security import create_access_token
from app.models import AssignedReviewer, Organisation, OrgAssignment, OrgMembership, User, utcnow
from tests.test_org_isolation import _doc, _h, _join, _org, _user, world  # noqa: F401  (fixture)


@pytest.fixture()
def team(db, world):
    turacoz = db.execute(select(Organisation).where(Organisation.kind == "system")).scalar_one()
    reviewer = _user(db, "trz_reviewer", user_type="reviewer")
    _join(db, turacoz, reviewer)
    return turacoz, reviewer


def _assign(client, login_as, org, user, role="reviewer", designation=""):
    return client.post(f"/api/superadmin/organisations/{org.id}/assignments", headers=login_as("admin"),
                       json={"user_id": user.id, "role": role, "designation": designation})


def test_system_user_sees_only_assigned_workspaces(client, login_as, world, team):
    _, reviewer = team
    h = _h(reviewer)
    me = client.get("/api/auth/me", headers=h).json()
    assert me["organisation"]["kind"] == "system" and me["is_org_admin"] is False
    assert client.get(f"/api/doccore/documents/{world['doc_b'].id}", headers=h).status_code in (403, 404)

    assert _assign(client, login_as, world["other"], reviewer).status_code == 200
    h = _h(reviewer)
    assert client.get(f"/api/doccore/documents/{world['doc_b'].id}", headers=h).status_code != 404
    assert client.get(f"/api/doccore/documents/{world['doc_a'].id}", headers=h).status_code in (403, 404)  # SMBT: not assigned


def test_workspace_admin_works_with_assigned_reviewer(client, db, login_as, world, team):
    _, reviewer = team
    other_admin = _h(world["b_admin"])
    form = client.get(f"/api/library/documents/{world['doc_b'].id}/assign-reviewer", headers=other_admin).json()
    assert reviewer.username not in str(form)  # not assigned yet: not offered
    _assign(client, login_as, world["other"], reviewer)
    form = client.get(f"/api/library/documents/{world['doc_b'].id}/assign-reviewer", headers=other_admin).json()
    assert any(r.get("id") == reviewer.id for r in form["reviewers"])
    r = client.post(f"/api/library/documents/{world['doc_b'].id}/assign-reviewer", headers=other_admin,
                    json={"reviewer_id": str(reviewer.id), "review_date": "2030-01-10", "remarks": ""})
    assert r.status_code == 200, r.text
    # the assignment belongs to the document's workspace, so its admin sees it
    manage = client.get(f"/api/library/documents/{world['doc_b'].id}/manage-reviewers", headers=other_admin)
    assert reviewer.username in manage.text


def test_points_of_contact(client, login_as, world, team):
    _, reviewer = team
    h = _h(world["b_admin"])
    r = client.post(f"/api/org/members/{world['b_member'].id}/poc", headers=h, json={"is_poc": True, "designation": "HoD Pharmacy"})
    assert r.status_code == 200, r.text
    assert client.post(f"/api/org/members/{world['a_member'].id}/poc", headers=h, json={"is_poc": True}).status_code == 404
    _assign(client, login_as, world["other"], reviewer, role="poc", designation="Account manager")
    c = client.get("/api/org/contacts", headers=h).json()
    assert [(p["username"], p["designation"]) for p in c["pocs"]] == [(world["b_member"].username, "HoD Pharmacy")]
    assert [(t["username"], t["role_label"], t["designation"]) for t in c["turacoz"]] == [(reviewer.username, "PoC", "Account manager")]
    # the superadmin sees the same on the organisation page and can remove the assignment
    detail = client.get(f"/api/superadmin/organisations/{world['other'].id}", headers=login_as("admin")).json()
    a_id = detail["contacts"]["turacoz"][0]["id"]
    assert client.delete(f"/api/superadmin/organisations/{world['other'].id}/assignments/{a_id}",
                         headers=login_as("admin")).status_code == 200
    assert client.get("/api/org/contacts", headers=h).json()["turacoz"] == []


def test_only_system_users_can_be_assigned(client, login_as, world):
    r = _assign(client, login_as, world["other"], world["a_member"])
    assert r.status_code == 400 and "not a Turacoz system user" in r.json()["detail"]
    found = client.get("/api/superadmin/organisations/system-users/search", headers=login_as("admin")).json()["users"]
    assert found and all(u["username"] != world["a_member"].username for u in found)


def test_existing_turacoz_reviewers_keep_their_smbt_reviews(client, db):
    """Data check: the reviewers moved to the Turacoz team still see their SMBT review assignments."""
    anuradha = db.execute(select(User).where(User.username == "anuradha.chatterjee")).scalar_one_or_none()
    if anuradha is None:
        pytest.skip("not in this database")
    total = db.execute(select(AssignedReviewer).where(AssignedReviewer.user_id == anuradha.id,
                                                      AssignedReviewer.is_active == 1)).scalars().all()
    h = {"Authorization": f"Bearer {create_access_token(anuradha.id)}"}
    d = client.get("/api/doccore/reviewer-dashboard", headers=h)
    assert d.status_code == 200, d.text
    assert d.json()["total_assignments"] == len(total)


def test_team_members_show_their_review_work(client, db, login_as, world, team):
    import datetime
    turacoz, reviewer = team
    from app.models import DocumentVersion
    v = db.execute(select(DocumentVersion).where(DocumentVersion.document_id == world["doc_b"].id)).scalar_one()
    now = utcnow()
    for done, due, started in ((1, now + datetime.timedelta(days=5), now), (0, now - datetime.timedelta(days=2), None),
                               (0, now + datetime.timedelta(days=3), now)):
        db.add(AssignedReviewer(user_id=reviewer.id, document_version_id=v.id, review_date=due, is_completed=done,
                                is_active=1, remarks="", created_by_id=reviewer.id, modified_by_id=reviewer.id,
                                started_at=started, completed_at=now if done else None))
    db.commit()
    _assign(client, login_as, world["other"], reviewer)
    members = client.get(f"/api/superadmin/organisations/{turacoz.id}", headers=login_as("admin")).json()["members"]
    row = next(m for m in members if m["user_id"] == reviewer.id)
    assert row["reviews"] == {"assigned": 3, "completed": 1, "pending": 2, "in_progress": 1, "overdue": 1}
    assert row["assigned_to"] == ["Other University (Reviewer)"]


def test_review_breakdown_detail_and_report(client, db, login_as, world, team):
    import datetime
    import io
    from openpyxl import load_workbook
    from app.models import DocumentVersion
    turacoz, reviewer = team
    now = utcnow()
    for doc, done in ((world["doc_b"], 1), (world["doc_b"], 0), (world["doc_a"], 0)):  # Other University ×2, SMBT ×1
        v = db.execute(select(DocumentVersion).where(DocumentVersion.document_id == doc.id)).scalar_one()
        db.add(AssignedReviewer(user_id=reviewer.id, document_version_id=v.id, review_date=now + datetime.timedelta(days=3),
                                is_completed=done, is_active=1, remarks="", created_by_id=reviewer.id,
                                modified_by_id=reviewer.id, completed_at=now if done else None))
    db.commit()
    h = login_as("admin")
    members = client.get(f"/api/superadmin/organisations/{turacoz.id}", headers=h).json()["members"]
    row = next(m for m in members if m["user_id"] == reviewer.id)
    per_org = {r["organisation"]: (r["assigned"], r["completed"], r["pending"]) for r in row["reviews_by_organisation"]}
    assert per_org == {"Other University": (2, 1, 1), "SMBT": (1, 0, 1)}

    d = client.get(f"/api/superadmin/organisations/{turacoz.id}/members/{reviewer.id}/reviews", headers=h).json()
    assert d["totals"]["assigned"] == 3 and len(d["items"]) == 3
    assert {i["organisation"] for i in d["items"]} == {"Other University", "SMBT"}
    assert client.get(f"/api/superadmin/organisations/{turacoz.id}/members/{world['a_member'].id}/reviews",
                      headers=h).status_code == 404

    x = client.get(f"/api/superadmin/organisations/{turacoz.id}/review-report", headers=h)
    assert x.status_code == 200 and "attachment" in x.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "By university", "Assignments"]
    by_uni = [r for r in wb["By university"].iter_rows(values_only=True) if r[1] == reviewer.username]
    assert sorted((r[2], r[3]) for r in by_uni) == [("Other University", 2), ("SMBT", 1)]
    c = client.get(f"/api/superadmin/organisations/{turacoz.id}/review-report", params={"format": "csv"}, headers=h)
    assert c.status_code == 200 and reviewer.username in c.content.decode("utf-8-sig")
    assert client.get(f"/api/superadmin/organisations/{turacoz.id}/review-report", headers=login_as("teacher1")).status_code == 403


def test_university_admin_sees_only_reviews_of_their_documents(client, db, login_as, world, team):
    import datetime
    import io
    from openpyxl import load_workbook
    from app.models import DocumentVersion
    _, reviewer = team
    now = utcnow()
    for doc, done in ((world["doc_b"], 1), (world["doc_b"], 0), (world["doc_a"], 0)):  # Other University ×2, SMBT ×1
        v = db.execute(select(DocumentVersion).where(DocumentVersion.document_id == doc.id)).scalar_one()
        db.add(AssignedReviewer(user_id=reviewer.id, document_version_id=v.id, review_date=now + datetime.timedelta(days=3),
                                is_completed=done, is_active=1, remarks="", created_by_id=reviewer.id,
                                modified_by_id=reviewer.id, completed_at=now if done else None))
    db.commit()
    _assign(client, login_as, world["other"], reviewer)
    h = _h(world["b_admin"])
    d = client.get("/api/org/reviewers", headers=h).json()
    row = next(r for r in d["reviewers"] if r["user_id"] == reviewer.id)
    assert (row["assigned"], row["completed"], row["pending"]) == (2, 1, 1)  # not the SMBT review
    assert row["source"] == "turacoz" and row["assigned_to_us"] is True
    detail = client.get(f"/api/org/reviewers/{reviewer.id}", headers=h).json()
    assert {i["organisation"] for i in detail["items"]} == {"Other University"} and len(detail["items"]) == 2
    assert client.get(f"/api/org/reviewers/{world['a_member'].id}", headers=h).status_code == 404

    x = client.get("/api/org/review-report", headers=h)
    wb = load_workbook(io.BytesIO(x.content))
    rows = [r for r in wb["Assignments"].iter_rows(values_only=True)][1:]
    assert len(rows) == 2 and all(r[2] == "Other University" for r in rows)
    assert client.get("/api/org/review-report", params={"format": "csv"}, headers=h).status_code == 200
    assert client.get("/api/org/reviewers", headers=_h(world["b_member"])).status_code == 403
